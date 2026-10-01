"""知识库摄取：解析(PDF/Word/MD/TXT) → 结构感知分块 → 嵌入 → 双路入库。

结构感知：MD/DOCX 按标题层级切、保留 section_path；PDF 按页 + 页内段落；
TXT 按空行段落。块长上限 600 字（硬约束放代码层）。
嵌入失败（本地模型挂）时文档标记 failed、可重试，不阻塞关键词路入库。
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import delete as sa_delete
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.db.models import KbChunk, KbDocument
from backend.rag import vecstore
from backend.rag.bm25 import bm25_index
from backend.rag.embedder import Embedder, EmbedderUnavailable

MAX_CHUNK_CHARS = 600
MIN_CHUNK_CHARS = 40


@dataclass
class ParsedChunk:
    content: str
    section_path: str
    page_no: int | None


def _split_long(text: str, limit: int = MAX_CHUNK_CHARS) -> list[str]:
    """超长段落按句号/换行二次切；无句读边界时按字硬切（块长上限是硬约束）。"""
    if len(text) <= limit:
        return [text]
    pieces: list[str] = []
    buf = ""
    for seg in re.split(r"(?<=[。！？；\n])", text):
        if len(buf) + len(seg) > limit and buf:
            pieces.append(buf.strip())
            buf = seg
        else:
            buf += seg
    if buf.strip():
        pieces.append(buf.strip())
    hard: list[str] = []
    for p in pieces:
        while len(p) > limit:
            hard.append(p[:limit])
            p = p[limit:]
        if p:
            hard.append(p)
    return hard


def parse_markdown(text: str, min_chars: int = MIN_CHUNK_CHARS) -> list[ParsedChunk]:
    chunks: list[ParsedChunk] = []
    headings: dict[int, str] = {}  # 各级标题 -> 当前值；同级替换、深层自动清空
    buf: list[str] = []

    def current_path() -> str:
        return " > ".join(headings[l] for l in sorted(headings))

    def flush() -> None:
        body = "\n".join(buf).strip()
        buf.clear()
        if len(body) < min_chars:
            return
        for piece in _split_long(body):
            chunks.append(ParsedChunk(piece, current_path() or "正文", None))

    for line in text.splitlines():
        m = re.match(r"^(#{1,4})\s+(.+)$", line)
        if m:
            flush()
            level = len(m.group(1))
            headings[level] = m.group(2).strip()
            for deeper in [lv for lv in headings if lv > level]:
                del headings[deeper]
        else:
            buf.append(line)
    flush()
    return chunks


def parse_plain(text: str, source_note: str = "", min_chars: int = MIN_CHUNK_CHARS) -> list[ParsedChunk]:
    chunks: list[ParsedChunk] = []
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    buf = ""
    for para in paras:
        if len(buf) + len(para) > MAX_CHUNK_CHARS and buf:
            chunks.append(ParsedChunk(buf.strip(), source_note or "正文", None))
            buf = ""
        buf += ("\n" if buf else "") + para
    if buf.strip() and len(buf.strip()) >= min_chars:
        chunks.append(ParsedChunk(buf.strip(), source_note or "正文", None))
    return (
        [ParsedChunk(piece, p.section_path, p.page_no) for p in chunks for piece in _split_long(p.content)]
        or chunks
    )


def parse_pdf_bytes(data: bytes) -> list[ParsedChunk]:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    chunks: list[ParsedChunk] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = re.sub(r"[ \t]+", " ", text)
        for para in [p.strip() for p in re.split(r"\n\s*\n", text) if len(p.strip()) >= MIN_CHUNK_CHARS]:
            for piece in _split_long(para):
                chunks.append(ParsedChunk(piece, f"第 {page_no} 页", page_no))
    return chunks


def parse_docx_bytes(data: bytes) -> list[ParsedChunk]:
    from io import BytesIO

    import docx  # python-docx

    document = docx.Document(BytesIO(data))
    chunks: list[ParsedChunk] = []
    path: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        buf.clear()
        if len(body) < MIN_CHUNK_CHARS:
            return
        for piece in _split_long(body):
            chunks.append(ParsedChunk(piece, " > ".join(path) or "正文", None))

    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        text = para.text.strip()
        if style.startswith("heading") and text:
            flush()
            try:
                level = int(style.replace("heading", "").strip() or "1")
            except ValueError:
                level = 1
            path = path[: level - 1] + [text]
        elif text:
            buf.append(text)
    flush()
    return chunks


def ingest_document(
    db: Session,
    *,
    title: str,
    category: str,
    filename: str = "",
    source_url: str = "",
    source_note: str = "",
    markdown_text: str | None = None,
    file_bytes: bytes | None = None,
    file_kind: str = "",
    embedder: Embedder | None = None,
    min_chunk_chars: int = MIN_CHUNK_CHARS,
) -> tuple[KbDocument, list[KbChunk]]:
    """解析并入库一个文档。返回 (文档, 块列表)。抛 EmbedderUnavailable 时文档记 failed。

    min_chunk_chars：案例等短文档传 1，避免最小块长把整篇过滤成 0 块。
    """
    suffix = Path(filename).suffix.lower() if filename else file_kind
    if markdown_text is not None:
        parsed = parse_markdown(markdown_text, min_chars=min_chunk_chars)
    elif suffix == ".pdf" and file_bytes is not None:
        parsed = parse_pdf_bytes(file_bytes)
    elif suffix in (".docx",) and file_bytes is not None:
        parsed = parse_docx_bytes(file_bytes)
    elif suffix in (".md", ".markdown") and file_bytes is not None:
        parsed = parse_markdown(file_bytes.decode("utf-8", errors="replace"))
    elif suffix in (".txt", ".json", ".csv") and file_bytes is not None:
        parsed = parse_plain(file_bytes.decode("utf-8", errors="replace"), source_note)
    else:
        raise ValueError(f"不支持的文件类型：{suffix or '未知'}")

    doc = KbDocument(
        title=title,
        filename=filename,
        category=category,
        source_url=source_url,
        source_note=source_note,
        status="parsing",
    )
    db.add(doc)
    db.flush()

    chunks: list[KbChunk] = []
    for seq, parsed_chunk in enumerate(parsed):
        chunk = KbChunk(
            document_id=doc.id,
            seq=seq,
            content=parsed_chunk.content,
            section_path=parsed_chunk.section_path,
            page_no=parsed_chunk.page_no,
            tokens=len(parsed_chunk.content),
        )
        db.add(chunk)
        chunks.append(chunk)
    db.flush()

    # 双路入库：关键词路必成（BM25 增量）；向量路尽力（失败则文档标 failed、可重试）
    for chunk in chunks:
        bm25_index.add(chunk.id, chunk.content, category)
    embedder = embedder or Embedder()
    try:
        if vecstore.ensure_vec_table(db):
            vectors = embedder.embed([c.content for c in chunks])
            for chunk, vec in zip(chunks, vectors):
                vecstore.upsert(db, chunk.id, vec)
    except EmbedderUnavailable:
        doc.status = "failed"
        db.commit()
        raise

    doc.status = "ready"
    doc.chunk_count = len(chunks)
    db.commit()
    return doc, chunks


def delete_document(db: Session, document_id: int) -> bool:
    doc = db.get(KbDocument, document_id)
    if doc is None:
        return False
    chunk_rows = db.execute(select(KbChunk.id).where(KbChunk.document_id == document_id)).all()
    chunk_ids = [int(r[0]) for r in chunk_rows]
    vecstore.delete(db, chunk_ids)
    bm25_index.remove(chunk_ids)
    db.execute(sa_delete(KbChunk).where(KbChunk.document_id == document_id))
    db.delete(doc)
    db.commit()
    return True


def export_chunks_json(chunks: list[KbChunk]) -> str:
    return json.dumps(
        [
            {"id": c.id, "seq": c.seq, "section_path": c.section_path, "content": c.content}
            for c in chunks
        ],
        ensure_ascii=False,
        indent=2,
    )
